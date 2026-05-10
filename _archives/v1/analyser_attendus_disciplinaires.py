# =============================================================================
# analyser_attendus_disciplinaires.py
# =============================================================================
# Pour chaque activite en base (avec fichier _DESC_) :
#   - Envoie nom + contenu _DESC_ a Gemini
#   - Gemini identifie les attendus disciplinaires pertinents
#   - Insere uniquement dans activite_attendu (ne touche pas aux autres champs)
#   - Met a jour activite_cycle si de nouveaux cycles sont couverts
#
# Prerequis :
#   - migration_ajout_type_attendu.py lance
#   - import_attendus_disciplinaires.py lance
#   pip install google-generativeai python-docx pillow --break-system-packages
# =============================================================================

import base64
import json
import sqlite3
import sys
import time
from pathlib import Path

# =============================================================================
# CONFIGURATION
# =============================================================================

DB_PATH = Path(__file__).parent / "activites.db"
DROPBOX_ACTIVITES = Path(r"C:\Users\moina\Dropbox\Animation\Activités v2")

GEMINI_MODEL = "gemini-2.5-flash"
DELAI_ENTRE_APPELS = 15  # secondes entre appels API

# None = toutes les activites, ex: 5 = mode test
LIMITE = None

# True = retraiter meme les activites qui ont deja des attendus disciplinaires
FORCER_REANALYSE = False

# Extensions supportees pour les fichiers _DESC_
EXTENSIONS_SUPPORTEES = {".pdf", ".jpg", ".jpeg", ".png", ".docx", ".txt"}

# =============================================================================
# SYSTEM PROMPT (identique a analyser_activites.py)
# =============================================================================

SYSTEM_PROMPT = """Tu es un assistant specialise en pedagogie de la nature et en education a l'environnement.
Tu analyses des fiches d'activites pedagogiques nature et tu extrais des informations structurees.
Tu reponds UNIQUEMENT en JSON valide, sans aucun texte autour, sans balises markdown.
"""

# =============================================================================
# BASE DE DONNEES
# =============================================================================

def get_db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON")
    return conn


def charger_attendus_disciplinaires(conn):
    rows = conn.execute("""
        SELECT a.id, a.cycle_id, c.code, a.domaine, a.sous_domaine, a.libelle
        FROM attendu_scolaire a
        JOIN cycle c ON c.id = a.cycle_id
        WHERE a.type = 'disciplinaire'
        ORDER BY a.cycle_id, a.domaine, a.id
    """).fetchall()
    return [dict(r) for r in rows]


def activites_a_traiter(conn, forcer=False):
    if forcer:
        rows = conn.execute("""
            SELECT id, nom, chemin_dossier, description, objectif_texte
            FROM activite
            ORDER BY nom
        """).fetchall()
    else:
        deja = conn.execute("""
            SELECT DISTINCT aa.activite_id
            FROM activite_attendu aa
            JOIN attendu_scolaire a ON a.id = aa.attendu_id
            WHERE a.type = 'disciplinaire'
        """).fetchall()
        deja_ids = {r[0] for r in deja}
        rows = conn.execute("""
            SELECT id, nom, chemin_dossier, description, objectif_texte
            FROM activite
            ORDER BY nom
        """).fetchall()
        rows = [r for r in rows if r["id"] not in deja_ids]
    return [dict(r) for r in rows]


def enregistrer_attendus(conn, activite_id, attendu_ids):
    liens_ajoutes = 0
    for aid in attendu_ids:
        conn.execute(
            "INSERT OR IGNORE INTO activite_attendu (activite_id, attendu_id) VALUES (?, ?)",
            (activite_id, aid)
        )
        if conn.execute("SELECT changes()").fetchone()[0] > 0:
            liens_ajoutes += 1

    if attendu_ids:
        placeholders = ",".join("?" * len(attendu_ids))
        cycle_ids = [r[0] for r in conn.execute(
            f"SELECT DISTINCT cycle_id FROM attendu_scolaire WHERE id IN ({placeholders})",
            attendu_ids
        ).fetchall()]
        for cid in cycle_ids:
            conn.execute(
                "INSERT OR IGNORE INTO activite_cycle (activite_id, cycle_id) VALUES (?, ?)",
                (activite_id, cid)
            )
    conn.commit()
    return liens_ajoutes


# =============================================================================
# LECTURE DES FICHIERS _DESC_ (identique a analyser_activites.py)
# =============================================================================

def normaliser_chemin(chemin_dossier):
    if chemin_dossier and chemin_dossier.startswith("file:///"):
        chemin_dossier = chemin_dossier[8:].replace("/", "\\")
    return chemin_dossier


def trouver_fichiers_desc(chemin_dossier):
    if not chemin_dossier:
        return []
    dossier = Path(normaliser_chemin(chemin_dossier))
    if not dossier.exists():
        return []
    fichiers = []
    for f in dossier.iterdir():
        if f.is_file() and f.name.startswith("_DESC_") and f.suffix.lower() in EXTENSIONS_SUPPORTEES:
            fichiers.append(f)
    return fichiers


def fichier_vers_contenu_api(fichier: Path):
    ext = fichier.suffix.lower()

    if ext in {".jpg", ".jpeg", ".png"}:
        media_type = "image/jpeg" if ext in {".jpg", ".jpeg"} else "image/png"
        data = base64.standard_b64encode(fichier.read_bytes()).decode("utf-8")
        return {
            "type": "image",
            "source": {"type": "base64", "media_type": media_type, "data": data}
        }

    if ext == ".pdf":
        data = base64.standard_b64encode(fichier.read_bytes()).decode("utf-8")
        return {
            "type": "document",
            "source": {"type": "base64", "media_type": "application/pdf", "data": data}
        }

    if ext == ".docx":
        try:
            from docx import Document
            doc = Document(str(fichier))
            texte = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        except Exception as e:
            texte = f"[Erreur lecture DOCX : {e}]"
        return {"type": "text", "text": f"[Contenu du fichier {fichier.name}]\n{texte}"}

    if ext == ".txt":
        try:
            texte = fichier.read_text(encoding="utf-8", errors="replace")
        except Exception as e:
            texte = f"[Erreur lecture TXT : {e}]"
        return {"type": "text", "text": f"[Contenu du fichier {fichier.name}]\n{texte}"}

    return None


# =============================================================================
# PROMPT (adapte depuis analyser_activites.py)
# =============================================================================

def construire_prompt(activite_nom, referentiel_attendus):

    def format_liste(items):
        lignes = []
        for item in items:
            dom = item["domaine"]
            if item["sous_domaine"]:
                dom += f" / {item['sous_domaine']}"
            lignes.append(f"  id={item['id']}, cycle={item['code']}, domaine={dom}, libelle={item['libelle']}")
        return "\n".join(lignes)

    ref_text = f"""
=== ATTENDUS SCOLAIRES DISCIPLINAIRES (hors EDD) ===
Ces attendus sont issus des programmes officiels (BO), par matiere disciplinaire.

{format_liste(referentiel_attendus)}
"""

    prompt = f"""Tu vas analyser la fiche descriptive de l'activite pedagogique nature intitulee : \u00ab {activite_nom} \u00bb.

{ref_text}

=== INSTRUCTIONS ===

En te basant UNIQUEMENT sur le contenu des fichiers fournis, identifie les attendus scolaires disciplinaires
directement travailles par cette activite et retourne-les en JSON :

{{
  "attendu_ids": [liste des id d'attendus scolaires disciplinaires correspondants]
}}

REGLES :
- Ne propose que des IDs qui existent dans le referentiel fourni ci-dessus.
- Si aucun attendu ne correspond, retourne un tableau vide [].
- Pour les attendus : sois tres selectif. Ne retiens que les attendus qui correspondent DIRECTEMENT et PRECISEMENT
  a ce que l'activite fait concretement -- pas a ce qu'elle pourrait theoriquement aborder.
  Un attendu ne doit etre retenu que si un animateur pourrait affirmer sans hesitation que l'activite le travaille.
  Si l'activite est simple et ciblee, 1 ou 2 attendus suffisent.
  Ne retiens jamais un attendu par defaut ou parce qu'il est vaguement lie au theme :
  en cas de doute, ne le retiens pas.
- Maximum 6 attendus par activite.
- Le cycle de l'attendu n'est pas un critere de filtrage : inclure tous les cycles pertinents.
- Le JSON doit etre valide et complet. Aucun texte en dehors du JSON.
"""
    return prompt


# =============================================================================
# APPEL GEMINI (identique a analyser_activites.py)
# =============================================================================

def _appel_gemini(client, content):
    import google.generativeai as genai

    model = genai.GenerativeModel(
        model_name=GEMINI_MODEL,
        system_instruction=SYSTEM_PROMPT
    )

    parts = []
    for bloc in content:
        if bloc["type"] == "text":
            parts.append(bloc["text"])
        elif bloc["type"] == "image":
            parts.append({
                "inline_data": {
                    "mime_type": bloc["source"]["media_type"],
                    "data": bloc["source"]["data"]
                }
            })
        elif bloc["type"] == "document":
            parts.append({
                "inline_data": {
                    "mime_type": "application/pdf",
                    "data": bloc["source"]["data"]
                }
            })

    response = model.generate_content(parts)
    return response.text.strip()


def analyser_activite(client, activite_nom, fichiers, referentiel_attendus):
    content = []
    content.append({"type": "text", "text": construire_prompt(activite_nom, referentiel_attendus)})

    for fichier in fichiers:
        bloc = fichier_vers_contenu_api(fichier)
        if bloc:
            content.append({"type": "text", "text": f"\n--- Fichier : {fichier.name} ---"})
            content.append(bloc)

    MAX_RETRIES = 5
    for tentative in range(MAX_RETRIES):
        try:
            raw = _appel_gemini(client, content)
            break
        except Exception as e:
            msg = str(e)
            if "rate_limit" in msg or "429" in msg or "quota" in msg.lower() or "Resource" in msg:
                attente = 60 * (tentative + 1)
                print(f"  Rate limit atteint. Attente {attente}s (tentative {tentative+1}/{MAX_RETRIES})...")
                time.sleep(attente)
                if tentative == MAX_RETRIES - 1:
                    raise
            else:
                raise

    try:
        data = json.loads(raw)
    except json.JSONDecodeError:
        if "```" in raw:
            raw = raw.split("```")[1]
            if raw.startswith("json"):
                raw = raw[4:]
        data = json.loads(raw.strip())

    # Si Gemini retourne une liste au lieu d'un objet, on normalise
    if isinstance(data, list):
        data = {"passages_cles": data, "attendu_ids": [p.get("attendu_id") for p in data if isinstance(p, dict) and "attendu_id" in p]}

    return data


# =============================================================================
# MAIN
# =============================================================================

def main():
    print("=" * 60)
    print("ANALYSE ATTENDUS DISCIPLINAIRES VIA GEMINI")
    print("=" * 60)

    if not DB_PATH.exists():
        print(f"ERREUR : Base SQLite introuvable : {DB_PATH}")
        print("   Lancez d'abord create_schema.py")
        sys.exit(1)

    import os
    gemini_key = os.environ.get("GEMINI_API_KEY")
    if not gemini_key:
        print("ERREUR : variable d'environnement GEMINI_API_KEY non definie.")
        sys.exit(1)

    import google.generativeai as genai
    genai.configure(api_key=gemini_key)
    client = genai
    print(f"  API : Google Gemini ({GEMINI_MODEL})")

    conn = get_db()

    cols = [row[1] for row in conn.execute("PRAGMA table_info(attendu_scolaire)").fetchall()]
    if "type" not in cols:
        print("ERREUR : colonne 'type' absente. Lancez d'abord migration_ajout_type_attendu.py")
        conn.close()
        sys.exit(1)

    print("\n Chargement des attendus disciplinaires...")
    referentiel_attendus = charger_attendus_disciplinaires(conn)
    if not referentiel_attendus:
        print("ERREUR : aucun attendu disciplinaire en base. Lancez d'abord import_attendus_disciplinaires.py")
        conn.close()
        sys.exit(1)
    print(f"  {len(referentiel_attendus)} attendus disciplinaires charges.")

    ids_valides = {a["id"] for a in referentiel_attendus}

    print("\n Recherche des activites a traiter...")
    activites = activites_a_traiter(conn, forcer=FORCER_REANALYSE)

    # Filtrer celles qui ont des fichiers _DESC_ (identique a analyser_activites.py)
    a_analyser = []
    sans_desc = []
    for act in activites:
        fichiers = trouver_fichiers_desc(act.get("chemin_dossier"))
        if fichiers:
            a_analyser.append((act, fichiers))
        else:
            sans_desc.append(act["nom"])

    print(f"  -> {len(a_analyser)} ont des fichiers _DESC_ a analyser")
    print(f"  -> {len(sans_desc)} n'ont pas de fichier _DESC_ (ignorees)")

    if LIMITE:
        a_analyser = a_analyser[:LIMITE]
        print(f"  Mode test : limite de {LIMITE} activite(s)")

    if not a_analyser:
        print("\nOK Rien a faire.")
        conn.close()
        return

    print(f"\n{'=' * 60}")
    print(f"Debut de l'analyse ({len(a_analyser)} activites)...")
    print(f"{'=' * 60}\n")

    succes = 0
    erreurs = 0

    for i, (act, fichiers) in enumerate(a_analyser, 1):
        nom = act["nom"]
        print(f"[{i}/{len(a_analyser)}] {nom}")
        print(f"   Fichiers : {', '.join(f.name for f in fichiers)}")

        try:
            resultats = analyser_activite(client, nom, fichiers, referentiel_attendus)
            attendu_ids = resultats.get("attendu_ids", [])
            if not isinstance(attendu_ids, list):
                attendu_ids = []
            attendu_ids = [aid for aid in attendu_ids if aid in ids_valides]

            passages = resultats.get("passages_cles", [])
            if passages:
                print("   Passages cles identifies : " + str(len(passages)))
                for p in passages:
                    extrait = str(p.get("extrait", ""))[:70]
                    print("     - " + repr(extrait))

            liens = enregistrer_attendus(conn, act["id"], attendu_ids)
            print(f"  OK Enregistre : {liens} attendu(s) ajoute(s)")
            if attendu_ids:
                print(f"    {len(attendu_ids)} attendu(s) associe(s)")
            succes += 1

        except Exception as e:
            print(f"  ERREUR : {e}")
            erreurs += 1

        if i < len(a_analyser):
            time.sleep(DELAI_ENTRE_APPELS)

        print()

    conn.close()

    print("=" * 60)
    print(f"Termine : {succes} succes, {erreurs} erreur(s)")
    print("=" * 60)


if __name__ == "__main__":
    main()
