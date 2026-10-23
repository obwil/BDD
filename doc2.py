# =============================================================================
# doc2.py
# =============================================================================
# « Déroulé et objectifs » (ex Doc 2) d'un dossier de séjour :
# « Séjour_JJ-MM_Nom_Déroulé&Objectifs-pédagogiques.docx ».
#
#   - remplir_doc2 : écrit les infos du Doc 3, les matinées (Doc 5) et les activités choisies.
#   - lire_doc2    : relit le document (tes retouches comprises) pour générer le nouveau document.
#
# Règle de sécurité : une case n'est modifiée que si elle est vide, ou si elle contient encore
# exactement ce que l'outil y avait écrit la dernière fois. Un texte que tu as retouché est conservé.
# =============================================================================

import re
import shutil
from pathlib import Path
from typing import Dict, List, Optional

import doc3 as D

JOURS_SEMAINE = ["Mardi", "Mercredi", "Jeudi", "Vendredi"]
JOURS_ACTIVITES = ["lundi", "mardi", "mercredi", "jeudi"]
VIDES = {"", "…", "...", "jj/mm/aaaa"}


def trouver_doc2(dossier: Path) -> Optional[Path]:
    for f in sorted(dossier.glob("Séjour_*Déroulé*.docx")):
        if not f.name.startswith("~$"):
            return f
    return None


def assurer_doc2(dossier: Path, s: dict) -> Path:
    """Retourne le document du dossier ; le copie depuis le modèle B s'il manque."""
    f = trouver_doc2(dossier)
    if f:
        return f
    modele = next(iter(sorted((D.dossier_annee() / D.NOM_MODELE_B).glob("Séjour_*Déroulé*.docx"))), None)
    if modele is None:
        raise FileNotFoundError(f"« déroulé et objectifs » introuvable dans le modèle B ({D.NOM_MODELE_B})")
    jj_mm = f"{s['date_debut'][8:10]}-{s['date_debut'][5:7]}"
    nom = modele.name.replace("00-00", jj_mm, 1).replace("_Nom_", f"_{D.nom_court(s.get('enseignant'))}_", 1)
    cible = dossier / nom
    shutil.copy2(modele, cible)
    return cible


def _fr(iso: str) -> str:
    return f"{iso[8:10]}/{iso[5:7]}/{iso[:4]}"


def _matins(texte: Optional[str]) -> Dict[str, str]:
    out = {}
    for l in (texte or "").splitlines():
        m = re.match(r"\s*(\w+)\s*:\s*(.+?)\s*$", l)
        if m:
            out[m.group(1).capitalize()] = m.group(2)
    return out


def _ecrire(cell, texte: str):
    """Écrit dans la cellule en gardant la mise en forme du premier morceau de texte."""
    p = cell.paragraphs[0]
    if p.runs:
        p.runs[0].text = texte
        for r in p.runs[1:]:
            r.text = ""
    else:
        p.add_run(texte)


def _style_comme(cible, modele):
    """Applique à la cellule le style de texte d'une cellule déjà rédigée (plus de gris italique du modèle)."""
    import copy
    ref = modele.paragraphs[0].runs[0] if modele.paragraphs[0].runs else None
    run = cible.paragraphs[0].runs[0] if cible.paragraphs[0].runs else None
    if ref is None or run is None:
        return
    ancien = run._r.rPr
    if ancien is not None:
        run._r.remove(ancien)
    if ref._r.rPr is not None:
        run._r.insert(0, copy.deepcopy(ref._r.rPr))


def _cellules_distinctes(row) -> list:
    vues, out = [], []
    for c in row.cells:
        if c._tc not in vues:
            vues.append(c._tc)
            out.append(c)
    return out


def _texte_detail(acts: List[dict], champ: str, prefixe: str = "") -> str:
    """Texte d'un créneau : celui de l'activité si elle est seule, sinon « titre : texte » pour chacune."""
    acts = [a for a in acts if a.get(champ)]
    if not acts:
        return ""
    if len(acts) == 1 and not prefixe:
        return acts[0][champ]
    return "\n".join(f"{prefixe}{a['nom']} : {a[champ]}" for a in acts)


def remplir_doc2(chemin: Path, s: dict, planning: Optional[dict] = None, ecrits: Optional[dict] = None) -> dict:
    """Met à jour le document. Retourne {'fichier', 'remplis', 'conserves', 'ecrits'} ;
    'ecrits' = ce que l'outil a écrit (à mémoriser pour la prochaine mise à jour)."""
    from docx import Document
    doc = Document(chemin)
    ecrits = dict(ecrits or {})
    nouveaux: Dict[str, str] = dict(ecrits)
    remplis: List[str] = []
    conserves: List[str] = []

    def poser(cell, valeur: Optional[str], cle: str, nom: str, style_de=None):
        if valeur is None:
            return
        actuel = cell.text.strip()
        avant = ecrits.get(cle)
        if actuel == valeur.strip():
            nouveaux[cle] = valeur
            return
        if actuel in VIDES or (avant is not None and actuel == avant.strip()):
            _ecrire(cell, valeur)
            if style_de is not None:
                _style_comme(cell, style_de)
            nouveaux[cle] = valeur
            remplis.append(nom)
        else:
            conserves.append(nom)

    # 1. « Séjour du jj/mm/aaaa au jj/mm/aaaa » (une seule fois : les dates ne changent pas)
    for p in doc.paragraphs:
        if p.text.strip().startswith("Séjour du"):
            dates = iter([_fr(s["date_debut"]), _fr(s["date_fin"])])
            fait = False
            for r in p.runs:
                if r.text.strip() == "jj/mm/aaaa":
                    r.text = next(dates)
                    r.italic = False
                    fait = True
            if fait:
                remplis.append("Dates du séjour")
            break

    # 2. Fiche du séjour (premier tableau)
    classe = " — ".join(x for x in (s.get("ecole"), s.get("niveau1")) if x)
    valeurs = {
        "thematique": s.get("thematique_excel"),
        "objectifsdusejour": s.get("objectifs_texte"),
        "ecoleclasse": classe or None,
        "enseignantreferent": s.get("enseignant"),
        "unite": (s.get("unite") or "").capitalize() or None,
    }
    for row in doc.tables[0].rows:
        cells = _cellules_distinctes(row)
        libelle = D.norm(cells[0].text)
        if libelle == "mediateurreferent":
            if len(cells) >= 5:
                poser(cells[2], s.get("mediateur_matin"), "fiche:mediateur_matin", "Médiateur du matin")
                poser(cells[4], D.mediateur(), "fiche:mediateur_am", "Médiateur de l'après-midi")
        elif libelle in valeurs and len(cells) >= 2:
            poser(cells[-1], valeurs[libelle], f"fiche:{libelle}", cells[0].text.strip())

    # 3. Emploi du temps (deuxième tableau)
    if len(doc.tables) > 1:
        edt = doc.tables[1]
        entetes = [c.text.strip().lower() for c in edt.rows[0].cells]
        lignes = {r.cells[0].text.strip().lower(): r for r in edt.rows[1:]}
        modele_style = lignes["matin"].cells[1] if "matin" in lignes else None

        matins = _matins(s.get("matin_texte"))
        if matins and "matin" in lignes:
            for i, jour in enumerate(entetes):
                if jour.capitalize() in JOURS_SEMAINE:
                    poser(lignes["matin"].cells[i], matins.get(jour.capitalize(), "—"),
                          f"edt:{jour}:matin", f"Matin du {jour}", modele_style)

        if planning is not None:
            for moment in ("après-midi", "veillée"):
                ligne = lignes.get(moment)
                if ligne is None:
                    continue
                for jour in JOURS_ACTIVITES:
                    if jour not in entetes:
                        continue
                    acts = planning[jour][moment]
                    poser(ligne.cells[entetes.index(jour)], " + ".join(a["nom"] for a in acts) or "…",
                          f"edt:{jour}:{moment}", f"{moment.capitalize()} du {jour}", modele_style)

    # 4. Détail des activités (troisième tableau)
    if planning is not None and len(doc.tables) > 2:
        for row in doc.tables[2].rows:
            cells = _cellules_distinctes(row)
            if len(cells) < 3:
                continue
            bloc = D.norm(cells[0].text)
            champ = "objectif_texte" if D.norm(cells[1].text).startswith("objectifs") else "description"
            court = "objectifs" if champ == "objectif_texte" else "descriptif"
            if bloc.startswith("veillee"):
                lignes_v = [t for j in JOURS_ACTIVITES
                            if (t := _texte_detail(planning[j]["veillée"], champ, prefixe=f"{j.capitalize()} — "))]
                poser(cells[2], "\n".join(lignes_v), f"detail:veillee:{court}", f"Veillée : {court}")
            else:
                for j in JOURS_ACTIVITES:
                    if bloc.startswith(j):
                        poser(cells[2], _texte_detail(planning[j]["après-midi"], champ), f"detail:{j}:{court}",
                              f"{j.capitalize()} après-midi : {court}")

    if remplis:
        doc.save(chemin)
    return {"fichier": str(chemin), "remplis": remplis, "conserves": conserves, "ecrits": nouveaux}


# ----------------------------------------------------------------------------
# Lecture : le nouveau document se base sur ce que contient le fichier
# ----------------------------------------------------------------------------

def lire_doc2(chemin: Path) -> dict:
    """Fiche et matinées telles qu'elles sont écrites dans le document (retouches comprises)."""
    from docx import Document
    doc = Document(chemin)
    out: dict = {"matins": {}}
    vide = lambda t: t.strip() in VIDES
    for row in doc.tables[0].rows:
        cells = _cellules_distinctes(row)
        lib = D.norm(cells[0].text)
        if lib == "mediateurreferent" and len(cells) >= 5:
            out["mediateur_matin"] = None if vide(cells[2].text) else cells[2].text.strip()
            out["mediateur_am"] = None if vide(cells[4].text) else cells[4].text.strip()
        elif len(cells) >= 2:
            cle = {"thematique": "thematique", "objectifsdusejour": "objectifs", "ecoleclasse": "ecole_classe",
                   "enseignantreferent": "enseignant", "unite": "unite"}.get(lib)
            if cle:
                out[cle] = None if vide(cells[-1].text) else cells[-1].text.strip()
    if len(doc.tables) > 1:
        edt = doc.tables[1]
        entetes = [c.text.strip().capitalize() for c in edt.rows[0].cells]
        ligne = next((r for r in edt.rows[1:] if r.cells[0].text.strip().lower() == "matin"), None)
        if ligne:
            for i, jour in enumerate(entetes):
                t = ligne.cells[i].text.strip()
                if jour in JOURS_SEMAINE and t not in VIDES and t != "—":
                    out["matins"][jour] = t
    return out
