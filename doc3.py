# =============================================================================
# doc3.py
# =============================================================================
# Lecture des documents de l'année scolaire (dossier A) :
#   - Doc 3 : « Projets enseignants Trimestre N.xlsx » (séjours du trimestre)
#   - Doc 5 : « roulement_activites_matin.xlsx » (matinées selon l'unité)
#   - Doc 6 : « Référentiel écoles.xlsx » (catégorie de chaque école)
# et gestion des dossiers de séjour (copie du modèle B).
# Fonctions pures : pas de FastAPI ici.
# =============================================================================

import json
import re
import shutil
import unicodedata
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Dict, List, Optional

BASE_DIR = Path(__file__).parent
CONFIG_PATH = BASE_DIR / "config.json"

DEFAULT_DOSSIER_ANNEE = r"C:\Users\moina\Dropbox\CPAC\Documents projets\2026-2027"
DEFAULT_MEDIATEUR = "Romain"
NOM_MODELE_B = "00000000 - Prénom Nom"
NOM_ROULEMENT = "roulement_activites_matin.xlsx"
NOM_REFERENTIEL = "Référentiel écoles.xlsx"

UNITES = {"R": "rose", "O": "orange", "V": "verte", "B": "bleue"}
JOURS_ROULEMENT = ["Mardi", "Mercredi", "Jeudi", "Vendredi"]


def norm(s) -> str:
    """Minuscules, sans accents ni ponctuation : pour comparer des écritures différentes."""
    s = unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]", "", s)


def _config() -> dict:
    if CONFIG_PATH.exists():
        try:
            return json.loads(CONFIG_PATH.read_text(encoding="utf-8"))
        except Exception:
            pass
    return {}


def dossier_annee() -> Path:
    return Path(_config().get("dossier_annee") or DEFAULT_DOSSIER_ANNEE)


def mediateur() -> str:
    return _config().get("mediateur") or DEFAULT_MEDIATEUR


def annee_debut() -> int:
    m = re.search(r"(20\d{2})-(20\d{2})", str(dossier_annee()))
    return int(m.group(1)) if m else date.today().year


def _txt(v) -> Optional[str]:
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def _date_scolaire(v) -> Optional[date]:
    """Les dates saisies sans année (04/01) sont lues avec la mauvaise année par Excel :
    on garde jour et mois, et on recalcule l'année scolaire (août à décembre = année de début)."""
    if isinstance(v, datetime):
        v = v.date()
    if not isinstance(v, date):
        return None
    a = annee_debut()
    return date(a if v.month >= 8 else a + 1, v.month, v.day)


# ----------------------------------------------------------------------------
# Doc 6 : référentiel des écoles
# ----------------------------------------------------------------------------

def lire_referentiel() -> Dict[str, str]:
    """{ écriture normalisée : catégorie } (le nom officiel et ses variantes)."""
    f = dossier_annee() / NOM_REFERENTIEL
    out: Dict[str, str] = {}
    if not f.exists():
        return out
    import openpyxl
    ws = openpyxl.load_workbook(f, data_only=True)["Écoles"]
    for row in ws.iter_rows(min_row=2, values_only=True):
        nom, cat, variantes = (row + (None, None, None))[:3]
        if not nom or not cat:
            continue
        out[norm(nom)] = str(cat).strip()
        for v in str(variantes or "").split(";"):
            if v.strip():
                out[norm(v)] = str(cat).strip()
    return out


# ----------------------------------------------------------------------------
# Doc 5 : roulement des matinées
# ----------------------------------------------------------------------------

def lire_roulement() -> Dict[str, Dict[str, str]]:
    """{ 'bleue': {'Mardi': 'Pôle biodiversité', ...}, ... }"""
    f = dossier_annee() / NOM_ROULEMENT
    out: Dict[str, Dict[str, str]] = {}
    if not f.exists():
        return out
    import openpyxl
    ws = openpyxl.load_workbook(f, data_only=True).worksheets[0]
    jours = {c.column: str(c.value).strip().capitalize() for c in ws[2] if c.value and c.column > 1}
    for row in ws.iter_rows(min_row=3):
        m = re.match(r"unit[ée]\s+(\w+)", str(row[0].value or "").strip(), re.I)
        if not m:
            continue
        out[norm(m.group(1))] = {jours[c.column]: str(c.value).strip() for c in row if c.column in jours and c.value}
    return out


def texte_matin(unite: Optional[str], date_debut: Optional[str], roulement: Dict[str, Dict[str, str]],
                date_fin: Optional[str] = None) -> Optional[str]:
    """« Mardi : Pôle biodiversité
Mercredi : … ». Semaine à 4 jours (séjour de 4 jours ou moins) :
    on supprime le mardi et on commence au mercredi. Les jours avant le début du séjour sont retirés."""
    jours = roulement.get(norm(unite))
    if not jours:
        return None
    debut = date.fromisoformat(date_debut) if date_debut else None
    fin = date.fromisoformat(date_fin) if date_fin else None
    quatre_jours = bool(debut and fin and (fin - debut).days <= 3)
    lundi = debut - timedelta(days=debut.weekday()) if debut else None
    lignes = []
    for i, j in enumerate(JOURS_ROULEMENT, start=1):
        if j not in jours or (quatre_jours and j == "Mardi"):
            continue
        if lundi and lundi + timedelta(days=i) < debut:
            continue
        lignes.append(f"{j} : {jours[j]}")
    return "\n".join(lignes) or None


# ----------------------------------------------------------------------------
# Doc 3 : Projets enseignants Trimestre N
# ----------------------------------------------------------------------------

COLONNES = {
    "dates": "dates", "ecoles": "ecole", "classes": "classes", "enseignants": "enseignants",
    "vtt": "vtt", "u": "unite", "mediateurmatin": "mediateur_matin",
    "mediateurapresmidi": "mediateur_am", "projets": "projet", "thematiques": "projet",
    "objectifspedagogiques": "objectifs", "delais": "delai",
}


def trouver_doc3(n: int):
    """Retourne (fichier, feuille) pour le trimestre n : d'abord dans le dossier du trimestre,
    sinon dans le dossier de l'année."""
    import openpyxl
    A = dossier_annee()
    nom = f"Projets enseignants Trimestre {n}.xlsx"
    for f in (A / f"Trimestre {n}" / nom, A / nom):
        if not f.exists():
            continue
        try:
            wb = openpyxl.load_workbook(f, data_only=True)
        except Exception:
            continue
        feuille = next((w for w in wb.worksheets if norm(w.title) == norm(f"Trimestre {n}")), None)
        if feuille is None and len(wb.worksheets) == 1:
            feuille = wb.worksheets[0]
        if feuille is not None:
            return f, feuille
    return None, None


def lire_doc3(n: int) -> dict:
    """Toutes les lignes du Doc 3 d'un trimestre, avec les dates de leur semaine.
    {'fichier': Path|None, 'lignes': [...], 'avertissements': [...]}"""
    f, ws = trouver_doc3(n)
    res = {"fichier": f, "lignes": [], "avertissements": []}
    if ws is None:
        return res

    # en-têtes
    ent_row, cols = None, {}
    for row in ws.iter_rows(min_row=1, max_row=12):
        trouve = {COLONNES[norm(c.value)]: c.column for c in row if c.value and norm(c.value) in COLONNES}
        if "ecole" in trouve and "enseignants" in trouve:
            ent_row, cols = row[0].row, trouve
            break
    if not cols:
        res["avertissements"].append(f"{f.name} : en-têtes introuvables")
        return res

    def val(r, k):
        return ws.cell(r, cols[k]).value if k in cols else None

    lignes = [r for r in range(ent_row + 1, ws.max_row + 1) if _txt(val(r, "ecole"))]
    # une semaine = 4 lignes consécutives (une par unité)
    for i in range(0, len(lignes), 4):
        bloc = lignes[i:i + 4]
        dates = [_date_scolaire(val(r, "dates")) for r in range(bloc[0], bloc[-1] + 1)]
        dates = [d for d in dates if d]
        delai = next((val(r, "delai") for r in bloc if val(r, "delai")), None)
        unites = [_txt(val(r, "unite")) for r in bloc]
        if len(dates) < 2 or len(bloc) != 4 or len({u for u in unites if u}) != len([u for u in unites if u]):
            res["avertissements"].append(
                f"{f.name} · lignes {bloc[0]}-{bloc[-1]} : semaine à vérifier (dates ou unités incohérentes)")
        if len(dates) < 2:
            continue
        for r in bloc:
            u = (_txt(val(r, "unite")) or "").upper()
            res["lignes"].append({
                "trimestre": n, "ligne": r, "fichier": f.name,
                "date_debut": dates[0].isoformat(), "date_fin": dates[1].isoformat(),
                "ecole": _txt(val(r, "ecole")), "classes": _txt(val(r, "classes")),
                "enseignants": _txt(val(r, "enseignants")),
                "unite_code": u, "unite": UNITES.get(u),
                "mediateur_matin": _txt(val(r, "mediateur_matin")),
                "mediateur_am": _txt(val(r, "mediateur_am")),
                "projet": _txt(val(r, "projet")), "objectifs": _txt(val(r, "objectifs")),
                "delai": (_date_delai(delai)),
            })
    return res


def _date_delai(v) -> Optional[str]:
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    return None


def mes_sejours(lignes: List[dict], role: str = "am") -> List[dict]:
    """Lignes où tu es médiateur d'après-midi (role='am') ou du matin (role='matin')."""
    cle = "mediateur_am" if role == "am" else "mediateur_matin"
    moi = norm(mediateur())
    return [l for l in lignes if norm(l[cle]) == moi]


# ----------------------------------------------------------------------------
# Noms et dossiers de séjour
# ----------------------------------------------------------------------------

def nom_court(enseignants: Optional[str]) -> str:
    """« Mme Braun » → « Braun » ; « Mmes Fumey et Kruba » → « Fumey et Kruba »."""
    s = (enseignants or "").strip()
    s = re.sub(r"^(?:(?:mmes?|mrs?|mm|m)\.?\s+)+", "", s, flags=re.I).strip()
    s = re.sub(r'[\\/:*?"<>|]', "", s).strip(" .")
    return s or "Inconnu"


def noms_famille(enseignants: Optional[str]) -> List[str]:
    return [n for n in (norm(x) for x in re.split(r"\s+et\s+|,", nom_court(enseignants))) if n]


def nom_dossier(s: dict) -> str:
    return f"{s['date_debut'].replace('-', '')} - {nom_court(s.get('enseignant'))}"


def dossier_trimestre(n: int) -> Path:
    return dossier_annee() / f"Trimestre {n}"


def trouver_dossier(s: dict) -> Optional[Path]:
    """Dossier existant du séjour : même date de début, et le nom de l'enseignant dedans."""
    base = dossier_trimestre(s["trimestre"]) if s.get("trimestre") else None
    zones = [base] if base and base.exists() else [p for p in dossier_annee().glob("Trimestre *") if p.is_dir()]
    prefixe = s["date_debut"].replace("-", "") + " - "
    noms = noms_famille(s.get("enseignant"))
    cands = [p for z in zones for p in z.iterdir() if p.is_dir() and p.name.startswith(prefixe)]
    for p in cands:
        if any(n in norm(p.name) for n in noms):
            return p
    return None


def creer_dossier(s: dict, dry_run: bool = False) -> dict:
    """Copie le dossier B dans le dossier du trimestre. Ne remplace jamais un dossier existant."""
    A = dossier_annee()
    modele = A / NOM_MODELE_B
    cible = dossier_trimestre(s["trimestre"]) / nom_dossier(s)
    if not modele.exists():
        return {"ok": False, "erreur": f"modèle introuvable : {modele}", "dossier": str(cible)}
    if cible.exists():
        return {"ok": False, "erreur": "existe déjà", "dossier": str(cible)}
    if not dry_run:
        shutil.copytree(modele, cible)
        jj_mm = f"{s['date_debut'][8:10]}-{s['date_debut'][5:7]}"
        for f in cible.iterdir():
            if f.name.startswith("Séjour_00-00_Nom_"):
                f.rename(f.with_name(f.name.replace("00-00", jj_mm, 1).replace("_Nom_", f"_{nom_court(s.get('enseignant'))}_", 1)))
    return {"ok": True, "dossier": str(cible)}
