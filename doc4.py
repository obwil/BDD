# =============================================================================
# doc4.py
# =============================================================================
# Doc 4 : « EDT trimestre N.xlsx » (dossier Pôle biodiversité de l'année).
# Pour chaque semaine où tu es médiateur du matin (Doc 3), 4 lignes : une par unité,
# dans l'ordre des jours où chaque unité vient au pôle biodiversité (Doc 5).
#
# Colonnes pré-remplies : Date, Jour, Unité, École, Catégorie, Niveau, Enseignant.
# Restent à toi : Effectif enfants, Effectif adultes, Activité.
# Une case déjà remplie n'est jamais modifiée.
# =============================================================================

import shutil
from copy import copy
from datetime import date, timedelta
from pathlib import Path
from typing import Dict, List, Optional

import doc3 as D

JOURS = ["Mardi", "Mercredi", "Jeudi", "Vendredi"]
COLONNES = ["Date", "Jour", "Unité", "Ecole", "Catégorie", "Niveau", "Enseignant",
            "Effectif enfants", "Effectif adultes", "Activité"]
LIGNES_PAR_SEMAINE = 4
PREMIERE_LIGNE = 2


def dossier_edt() -> Path:
    cfg = D._config().get("dossier_edt")
    if cfg:
        return Path(cfg)
    A = D.dossier_annee()  # …/CPAC/Documents projets/2026-2027 → …/CPAC/Pôle biodiversité/2026-2027
    return A.parent.parent / "Pôle biodiversité" / A.name


def fichier_edt(n: int) -> Path:
    return dossier_edt() / f"EDT trimestre {n}.xlsx"


def mode_unites() -> str:
    """'toutes' : les 4 unités de chaque semaine ; 'miennes' : seulement celles dont tu es médiateur du matin."""
    return D._config().get("doc4_unites") or "toutes"


def _fr(d: date) -> str:
    return f"{d.day:02d}/{d.month:02d}"


# ----------------------------------------------------------------------------
# Semaines et lignes à écrire
# ----------------------------------------------------------------------------

def semaines_matin(lignes: List[dict]) -> List[dict]:
    """Semaines (date de début, date de fin) où tu es médiateur du matin d'au moins une unité."""
    moi = D.norm(D.mediateur())
    par_semaine: Dict[tuple, List[dict]] = {}
    for l in lignes:
        par_semaine.setdefault((l["date_debut"], l["date_fin"]), []).append(l)
    out = []
    for (debut, fin), ls in sorted(par_semaine.items()):
        if any(D.norm(l["mediateur_matin"]) == moi for l in ls):
            out.append({"date_debut": debut, "date_fin": fin, "lignes": ls})
    return out


def jour_du_pole(roulement: Dict[str, Dict[str, str]]) -> List[tuple]:
    """[(unité, jour)] triés par jour : le jour où chaque unité est au pôle biodiversité."""
    out = []
    for unite, jours in roulement.items():
        jour = next((j for j, label in jours.items() if D.norm(label) == "polebiodiversite"), None)
        if jour:
            out.append((unite, jour))
    return sorted(out, key=lambda x: JOURS.index(x[1]) if x[1] in JOURS else 9)


def lignes_semaine(sem: dict, roulement: dict, referentiel: dict) -> List[dict]:
    """Les 4 lignes d'une semaine : {colonne: valeur}."""
    debut = date.fromisoformat(sem["date_debut"])
    fin = date.fromisoformat(sem["date_fin"])
    lundi = debut - timedelta(days=debut.weekday())
    quatre_jours = (fin - debut).days <= 3
    moi = D.norm(D.mediateur())
    par_unite = {D.norm(l["unite"]): l for l in sem["lignes"]}
    label = f"{_fr(debut)} au {_fr(fin)}"
    rows = []
    for unite, jour in jour_du_pole(roulement):
        l = par_unite.get(D.norm(unite))
        pas_de_pole = quatre_jours and jour == "Mardi"  # semaine à 4 jours : pas de mardi
        row = {"Date": label, "Unité": unite.capitalize()}
        row["Jour"] = "—" if pas_de_pole else f"{jour} {_fr(lundi + timedelta(days=JOURS.index(jour) + 1))}"
        if l and (mode_unites() == "toutes" or D.norm(l["mediateur_matin"]) == moi):
            row["Ecole"] = l["ecole"]
            row["Catégorie"] = referentiel.get(D.norm(l["ecole"]))
            row["Niveau"] = l["classes"]
            row["Enseignant"] = l["enseignants"]
        if pas_de_pole:
            row["Activité"] = "Pas de pôle biodiversité (semaine à 4 jours)"
        rows.append(row)
    return rows


# ----------------------------------------------------------------------------
# Écriture dans le classeur
# ----------------------------------------------------------------------------

def _vide(v) -> bool:
    return v is None or str(v).strip() in ("", "au")


def _copier_bloc_modele(ws, k: int):
    """Ajoute le bloc de semaine n°k en copiant la mise en forme du premier bloc."""
    r0 = PREMIERE_LIGNE + LIGNES_PAR_SEMAINE * k
    for i in range(LIGNES_PAR_SEMAINE):
        for c in range(1, len(COLONNES) + 1):
            src, dst = ws.cell(PREMIERE_LIGNE + i, c), ws.cell(r0 + i, c)
            if src.has_style:
                dst._style = copy(src._style)
        if ws.row_dimensions[PREMIERE_LIGNE + i].height:
            ws.row_dimensions[r0 + i].height = ws.row_dimensions[PREMIERE_LIGNE + i].height
    ws.merge_cells(start_row=r0, start_column=1, end_row=r0 + LIGNES_PAR_SEMAINE - 1, end_column=1)


def remplir_edt(n: int, lignes_doc3: List[dict], dry_run: bool = False) -> dict:
    """Écrit les semaines du trimestre n dans « EDT trimestre N.xlsx ».
    Retourne {'fichier', 'ok', 'semaines': [...], 'ecrites', 'conservees', 'blocs_ajoutes', 'erreur'}."""
    import openpyxl
    f = fichier_edt(n)
    res = {"fichier": str(f), "ok": False, "semaines": [], "ecrites": 0, "conservees": [], "blocs_ajoutes": 0}
    if not f.exists():
        res["erreur"] = f"Fichier introuvable : {f}"
        return res
    roulement, referentiel = D.lire_roulement(), D.lire_referentiel()
    if not roulement:
        res["erreur"] = "Roulement des matinées introuvable (Doc 5)."
        return res
    semaines = semaines_matin(lignes_doc3)
    try:
        wb = openpyxl.load_workbook(f)
    except PermissionError:
        res["erreur"] = "Le fichier est ouvert dans Excel : ferme-le puis réessaie."
        return res
    ws = wb.worksheets[0]
    nb_blocs = max(0, (ws.max_row - PREMIERE_LIGNE + 1) // LIGNES_PAR_SEMAINE)

    pris: set = set()  # blocs attribués pendant cet appel (nécessaire pour l'aperçu, qui n'écrit rien)

    def bloc_de(label: str) -> Optional[int]:
        for k in range(nb_blocs):
            if k not in pris and str(ws.cell(PREMIERE_LIGNE + LIGNES_PAR_SEMAINE * k, 1).value or "").strip() == label:
                return k
        return None

    def bloc_libre() -> Optional[int]:
        return next((k for k in range(nb_blocs)
                     if k not in pris and _vide(ws.cell(PREMIERE_LIGNE + LIGNES_PAR_SEMAINE * k, 1).value)), None)

    if not dry_run and any(not _vide(ws.cell(r, c).value) for r in range(2, ws.max_row + 1) for c in (8, 9, 10)):
        shutil.copy2(f, f.with_name(f.name + ".bak"))  # tes effectifs et activités : on garde une copie

    for sem in semaines:
        rows = lignes_semaine(sem, roulement, referentiel)
        label = rows[0]["Date"]
        k = bloc_de(label)
        if k is None:
            k = bloc_libre()
        if k is None:  # plus de bloc libre : on en ajoute un
            k = nb_blocs
            if not dry_run:
                _copier_bloc_modele(ws, k)
            nb_blocs += 1
            res["blocs_ajoutes"] += 1
        pris.add(k)
        ecrites = 0
        for i, row in enumerate(rows):
            r = PREMIERE_LIGNE + LIGNES_PAR_SEMAINE * k + i
            for nom, valeur in row.items():
                if valeur is None or (nom == "Date" and i > 0):  # la date est dans la cellule fusionnée du haut
                    continue
                cell = ws.cell(r, COLONNES.index(nom) + 1)
                actuel = cell.value
                if str(actuel or "").strip() == str(valeur).strip():
                    continue
                # un simple nom de jour (« Mardi ») est le texte du modèle, pas une saisie
                if _vide(actuel) or (nom == "Jour" and str(actuel).strip() in JOURS):
                    if not dry_run:
                        cell.value = valeur
                    ecrites += 1
                else:
                    res["conservees"].append(f"{label} · {nom} ({cell.coordinate})")
        res["semaines"].append({"semaine": label, "bloc": k + 1, "cases": ecrites})
        res["ecrites"] += ecrites

    if not dry_run:
        try:
            wb.save(f)
        except PermissionError:
            res["erreur"] = "Le fichier est ouvert dans Excel : ferme-le puis réessaie."
            return res
    res["ok"] = True
    return res
