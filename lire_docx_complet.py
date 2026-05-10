import os
import docx

# Lire les docx en entier
fichiers = [
    (226, r"C:\Users\moina\Dropbox\Animation\Activités v2\Concours d'anecdotes\_DESC_Concours d'anecdotes.docx"),
    (293, r"C:\Users\moina\Dropbox\Animation\Activités v2\Jeu des 10s\_DESC_Jeu des 10s.docx"),
    (486, r"C:\Users\moina\Dropbox\Animation\Activités v2\Histoire loopstation\_DESC_Histoire loopstation.docx"),
    (259, r"C:\Users\moina\Dropbox\Animation\Activités v2\Qu'est-ce qu'un être vivant\_DESC_Qu'est-ce qu'un être vivant.docx"),
    (921, r"C:\Users\moina\Dropbox\Animation\Activités v2\Balade musicale contée\_DESC_Balade musicale contée v3.docx"),
]

for aid, path in fichiers:
    print(f"\n{'='*60}")
    print(f"ID: {aid} | {os.path.basename(os.path.dirname(path))}")
    try:
        doc = docx.Document(path)
        # Tables aussi
        texte_paras = '\n'.join(p.text for p in doc.paragraphs if p.text.strip())
        texte_tables = ''
        for table in doc.tables:
            for row in table.rows:
                row_text = ' | '.join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    texte_tables += row_text + '\n'
        print(texte_paras[:4000])
        if texte_tables:
            print("\n[TABLES]")
            print(texte_tables[:2000])
    except Exception as e:
        print(f"  ERREUR: {e}")
