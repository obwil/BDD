# =============================================================================
# sejours.py
# =============================================================================
# Planification des séjours : lien avec le classeur Excel des plannings,
# association d'activités aux après-midi / veillées (lundi→jeudi) et
# génération d'un document Word récapitulatif par séjour.
#
# Branché dans main.py via : app.include_router(sejours.router)
# =============================================================================

import json
import os
import re
import sqlite3
from datetime import date
from pathlib import Path
from typing import List, Optional

from fastapi import APIRouter, HTTPException
from pydantic import BaseModel

import doc2 as D2
import doc3 as D

BASE_DIR = Path(__file__).parent
DB_PATH = BASE_DIR / "activites.db"
CONFIG_PATH = BASE_DIR / "config.json"
SORTIES_DIR = BASE_DIR / "sorties_sejours"

DEFAULT_EXCEL = r"C:\Users\moina\Dropbox\CPAC\Documents projets\2026-2027\Plannings séjours_2026-2027.xlsx"

JOURS = ["lundi", "mardi", "mercredi", "jeudi"]
MOMENTS = ["après-midi", "veillée"]

router = APIRouter(prefix="/api/planning", tags=["planning séjours"])


# ----------------------------------------------------------------------------
# Base de données
# ----------------------------------------------------------------------------

def get_db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON")
    return conn


# Colonnes ajoutées à la table `sejour` existante (reprise de l'Excel)
SEJOUR_COLS = {
    "date_code": "TEXT",      # ex "28090210" tel que saisi dans l'Excel
    "niveau1": "TEXT",        # col A
    "niveau2": "TEXT",        # col B
    "categorie": "TEXT",      # col C (Hors EP / REP / REP+)
    "ecole": "TEXT",          # col D
    "enseignant": "TEXT",     # col E
    "classe_ar": "TEXT",      # col H
    "unite": "TEXT",          # col I
    "thematique_excel": "TEXT",  # col J
    "matin_texte": "TEXT",    # cellule du haut de la col K (activités du matin)
    "excel_row": "INTEGER",
    # Doc 3 « Projets enseignants Trimestre N »
    "trimestre": "INTEGER",
    "doc3_key": "TEXT",           # « date_debut|code unité » : identifie le séjour d'une semaine
    "mediateur_matin": "TEXT",
    "delai": "TEXT",              # date limite du dossier (colonne Délais)
    "objectifs_texte": "TEXT",    # objectifs pédagogiques du Doc 3
    "dossier_path": "TEXT",       # dossier du séjour dans Dropbox
}


def ensure_schema():
    conn = get_db()
    existing = {r["name"] for r in conn.execute("PRAGMA table_info(sejour)")}
    for col, typ in SEJOUR_COLS.items():
        if col not in existing:
            conn.execute(f"ALTER TABLE sejour ADD COLUMN {col} {typ}")
    conn.execute("""
        CREATE TABLE IF NOT EXISTS sequence_activite_attendu (
            sequence_id INTEGER NOT NULL,
            activite_id INTEGER NOT NULL,
            attendu_id  INTEGER NOT NULL REFERENCES attendu_scolaire(id) ON DELETE CASCADE,
            PRIMARY KEY (sequence_id, activite_id, attendu_id),
            FOREIGN KEY (sequence_id, activite_id)
                REFERENCES sequence_activite(sequence_id, activite_id) ON DELETE CASCADE
        )
    """)
    conn.execute("""
        CREATE TABLE IF NOT EXISTS doc2_ecrit (
            sejour_id INTEGER NOT NULL REFERENCES sejour(id) ON DELETE CASCADE,
            cle       TEXT NOT NULL,
            texte     TEXT,
            PRIMARY KEY (sejour_id, cle)
        )
    """)
    conn.execute("""
        CREATE TABLE IF NOT EXISTS sejour_ressource (
            sejour_id   INTEGER NOT NULL REFERENCES sejour(id) ON DELETE CASCADE,
            activite_id INTEGER NOT NULL REFERENCES activite(id) ON DELETE CASCADE,
            ordre       INTEGER DEFAULT 0,
            PRIMARY KEY (sejour_id, activite_id)
        )
    """)
    conn.commit()
    conn.close()


def _label(s: dict) -> str:
    d1, d2 = s["date_debut"], s["date_fin"]
    plage = f"{d1[8:10]}/{d1[5:7]} → {d2[8:10]}/{d2[5:7]}"
    qui = s.get("enseignant") or s.get("ecole") or "?"
    niv = "/".join(x for x in (s.get("niveau1"), s.get("niveau2")) if x)
    return f"{plage} · {qui}" + (f" ({niv})" if niv else "")


# ----------------------------------------------------------------------------
# Synchronisation Doc 3 → BDD
# ----------------------------------------------------------------------------

def _sejour_depuis_ligne(l: dict, referentiel: dict, roulement: dict) -> dict:
    s = {
        "date_debut": l["date_debut"], "date_fin": l["date_fin"],
        "niveau1": l["classes"], "niveau2": None, "public": l["classes"],
        "categorie": referentiel.get(D.norm(l["ecole"])),
        "ecole": l["ecole"], "enseignant": l["enseignants"], "unite": l["unite"],
        "thematique_excel": l["projet"], "thematique_generale": l["projet"],
        "objectifs_texte": l["objectifs"], "mediateur_matin": l["mediateur_matin"],
        "delai": l["delai"], "trimestre": l["trimestre"],
        "doc3_key": f"{l['date_debut']}|{l['unite_code']}",
        "matin_texte": D.texte_matin(l["unite"], l["date_debut"], roulement, l["date_fin"]),
        "excel_row": l["ligne"],
    }
    s["nom"] = _label(s)
    return s


COLS_MAJ = ["nom", "date_debut", "date_fin", "public", "niveau1", "niveau2", "categorie", "ecole",
            "enseignant", "unite", "thematique_excel", "thematique_generale", "objectifs_texte",
            "mediateur_matin", "delai", "trimestre", "doc3_key", "matin_texte", "excel_row"]


def _chercher_sejour(conn, s: dict):
    """Séjour déjà en base : même clé Doc 3, sinon un ancien séjour (import de l'ancien Excel)
    de même date de début dont l'enseignant correspond : il garde ses activités."""
    row = conn.execute("SELECT id FROM sejour WHERE doc3_key=?", (s["doc3_key"],)).fetchone()
    if row:
        return row["id"]
    noms = D.noms_famille(s["enseignant"])
    for r in conn.execute("SELECT id, enseignant FROM sejour WHERE doc3_key IS NULL AND date_debut=?",
                          (s["date_debut"],)):
        if any(n in D.norm(r["enseignant"]) for n in noms):
            return r["id"]
    return None


@router.post("/sync")
def synchroniser():
    """Lit les « Projets enseignants Trimestre N » et enregistre les séjours où tu es médiateur d'après-midi."""
    ensure_schema()
    referentiel, roulement = D.lire_referentiel(), D.lire_roulement()
    conn = get_db()
    if not conn.execute("SELECT 1 FROM sejour WHERE doc3_key IS NOT NULL").fetchone():
        import shutil  # première synchronisation : on garde une copie de la base
        shutil.copy2(DB_PATH, DB_PATH.with_name(DB_PATH.name + ".bak-avant-doc3"))
    ajoutes = maj = 0
    trimestres, avertissements = [], []
    if not roulement:
        avertissements.append("Roulement des matinées introuvable : les matinées ne seront pas renseignées.")
    for n in (1, 2, 3):
        doc = D.lire_doc3(n)
        avertissements += doc["avertissements"]
        if doc["fichier"] is None:
            trimestres.append({"trimestre": n, "fichier": None, "sejours": 0})
            continue
        mes = D.mes_sejours(doc["lignes"], "am")
        trimestres.append({"trimestre": n, "fichier": str(doc["fichier"]), "sejours": len(mes)})
        for l in mes:
            s = _sejour_depuis_ligne(l, referentiel, roulement)
            if not s["categorie"]:
                avertissements.append(f"École « {s['ecole']} » absente du référentiel des écoles.")
            if s["delai"] and s["delai"] > s["date_debut"]:
                avertissements.append(f"Délai du dossier ({s['delai']}) postérieur au séjour du {s['date_debut']} "
                                      f"({s['ecole']}, {s['enseignant']}) : à vérifier dans le Doc 3.")
            sid = _chercher_sejour(conn, s)
            if sid:
                conn.execute("UPDATE sejour SET " + ", ".join(f"{c}=:{c}" for c in COLS_MAJ) + " WHERE id=:id",
                             {**s, "id": sid})
                maj += 1
            else:
                sid = conn.execute("INSERT INTO sejour (" + ", ".join(COLS_MAJ) + ") VALUES ("
                                   + ", ".join(f":{c}" for c in COLS_MAJ) + ")", s).lastrowid
                ajoutes += 1
            d = D.trouver_dossier(s)
            conn.execute("UPDATE sejour SET dossier_path=? WHERE id=?", (str(d) if d else None, sid))
    orphelins = conn.execute("""
        SELECT s.nom FROM sejour s WHERE s.doc3_key IS NULL AND s.date_code IS NOT NULL AND EXISTS
          (SELECT 1 FROM sequence q JOIN sequence_activite sa ON sa.sequence_id=q.id WHERE q.sejour_id=s.id)
    """).fetchall()
    for o in orphelins:
        avertissements.append(f"Ancien séjour non retrouvé dans le Doc 3, avec des activités placées : {o['nom']}")
    conn.commit()
    manquants = _dossiers_manquants(conn)
    conn.close()
    return {"ajoutes": ajoutes, "mis_a_jour": maj, "trimestres": trimestres,
            "avertissements": sorted(set(avertissements)), "dossiers_manquants": len(manquants)}


# ----------------------------------------------------------------------------
# Lecture
# ----------------------------------------------------------------------------

@router.get("/sejours")
def liste_sejours():
    ensure_schema()
    conn = get_db()
    rows = conn.execute("""
        SELECT s.id, s.nom, s.date_debut, s.date_fin, s.enseignant, s.ecole, s.public, s.trimestre,
          s.categorie, s.delai, s.dossier_path,
          (SELECT COUNT(*) FROM sequence q JOIN sequence_activite sa ON sa.sequence_id=q.id
            WHERE q.sejour_id=s.id) AS nb_activites
        FROM sejour s WHERE s.doc3_key IS NOT NULL
        ORDER BY s.date_debut, s.excel_row""").fetchall()
    conn.close()
    out = []
    for r in rows:
        d = dict(r)
        d["dossier_existe"] = bool(d["dossier_path"] and Path(d["dossier_path"]).is_dir())
        out.append(d)
    return out


# ----------------------------------------------------------------------------
# Dossiers de séjour (copie du modèle B)
# ----------------------------------------------------------------------------

def _dossiers_manquants(conn) -> list:
    out = []
    for r in conn.execute("SELECT * FROM sejour WHERE doc3_key IS NOT NULL ORDER BY date_debut"):
        s = dict(r)
        if not (s["dossier_path"] and Path(s["dossier_path"]).is_dir()):
            out.append(s)
    return out


@router.get("/dossiers/manquants")
def dossiers_manquants():
    """Aperçu : dossiers qui seraient créés (rien n'est écrit)."""
    ensure_schema()
    conn = get_db()
    res = []
    for s in _dossiers_manquants(conn):
        r = D.creer_dossier(s, dry_run=True)
        res.append({"id": s["id"], "nom": s["nom"], **r})
    conn.close()
    return res


@router.post("/dossiers/creer")
def dossiers_creer():
    """Crée les dossiers manquants à partir du modèle B."""
    ensure_schema()
    conn = get_db()
    res = []
    for s in _dossiers_manquants(conn):
        r = D.creer_dossier(s)
        if r["ok"]:
            conn.execute("UPDATE sejour SET dossier_path=? WHERE id=?", (r["dossier"], s["id"]))
            conn.commit()
            r["doc2"] = _maj_doc2(conn, s["id"])
        res.append({"id": s["id"], "nom": s["nom"], **r})
    conn.commit()
    conn.close()
    return res


def _maj_doc2(conn, sejour_id: int, strict: bool = False) -> dict:
    """Écrit les infos du séjour et les activités choisies dans « déroulé et objectifs ».
    strict=True : lève une erreur HTTP (bouton) ; sinon renvoie un statut, sans jamais bloquer l'appelant."""
    def echec(code, message):
        if strict:
            raise HTTPException(code, message)
        return {"ok": False, "message": message}

    s = _charger_sejour(conn, sejour_id)
    dossier = Path(s["dossier_path"]) if s.get("dossier_path") and Path(s["dossier_path"]).is_dir() else D.trouver_dossier(s)
    if not dossier:
        return echec(409, "Pas de dossier pour ce séjour : crée-le d'abord.")
    try:
        fichier = D2.assurer_doc2(dossier, s)
        ecrits = {r["cle"]: r["texte"] for r in conn.execute(
            "SELECT cle, texte FROM doc2_ecrit WHERE sejour_id=?", (sejour_id,))}
        res = D2.remplir_doc2(fichier, s, s["planning"], ecrits)
    except PermissionError:
        return echec(423, "« Déroulé et objectifs » est ouvert dans Word : ferme-le puis réessaie.")
    except FileNotFoundError as e:
        return echec(404, str(e))
    conn.execute("DELETE FROM doc2_ecrit WHERE sejour_id=?", (sejour_id,))
    conn.executemany("INSERT INTO doc2_ecrit (sejour_id, cle, texte) VALUES (?,?,?)",
                     [(sejour_id, k, v) for k, v in res["ecrits"].items()])
    conn.commit()
    return {"ok": True, "fichier": res["fichier"], "remplis": res["remplis"], "conserves": res["conserves"]}


@router.post("/sejours/{sejour_id}/doc2")
def mettre_a_jour_doc2(sejour_id: int):
    """Met à jour « déroulé et objectifs » du dossier (le copie depuis le modèle B s'il manque).
    Les cases que tu as retouchées sont conservées."""
    ensure_schema()
    conn = get_db()
    if not conn.execute("SELECT 1 FROM sejour WHERE id=?", (sejour_id,)).fetchone():
        conn.close()
        raise HTTPException(404, "Séjour introuvable")
    res = _maj_doc2(conn, sejour_id, strict=True)
    conn.close()
    return res


def _charger_sejour(conn, sejour_id: int) -> dict:
    s = conn.execute("SELECT * FROM sejour WHERE id=?", (sejour_id,)).fetchone()
    if not s:
        raise HTTPException(404, "Séjour introuvable")
    res = dict(s)
    planning = {j: {m: [] for m in MOMENTS} for j in JOURS}
    seqs = conn.execute("SELECT * FROM sequence WHERE sejour_id=?", (sejour_id,)).fetchall()
    for q in seqs:
        if q["date"] not in planning or q["moment"] not in MOMENTS:
            continue
        acts = conn.execute("""
            SELECT a.id, a.nom, a.description, a.objectif_texte, sa.ordre
            FROM sequence_activite sa JOIN activite a ON a.id=sa.activite_id
            WHERE sa.sequence_id=? ORDER BY sa.ordre""", (q["id"],)).fetchall()
        for a in acts:
            d = dict(a)
            d["attendus"] = [dict(x) for x in conn.execute("""
                SELECT ats.id, ats.libelle, ats.domaine, ats.sous_domaine, ats.type, c.code AS cycle_code
                FROM sequence_activite_attendu saa
                JOIN attendu_scolaire ats ON ats.id=saa.attendu_id
                JOIN cycle c ON c.id=ats.cycle_id
                WHERE saa.sequence_id=? AND saa.activite_id=?
                ORDER BY c.id, ats.domaine""", (q["id"], a["id"]))]
            planning[q["date"]][q["moment"]].append(d)
    res["planning"] = planning
    res["ressources"] = [dict(r) for r in conn.execute("""
        SELECT a.id, a.nom, a.description, a.duree_min, a.effectif_min, a.effectif_max,
               a.lieu, a.format_groupe, sr.ordre
        FROM sejour_ressource sr JOIN activite a ON a.id = sr.activite_id
        WHERE sr.sejour_id = ? ORDER BY sr.ordre""", (sejour_id,))]
    return res


@router.get("/sejours/{sejour_id}")
def detail_sejour(sejour_id: int):
    ensure_schema()
    conn = get_db()
    res = _charger_sejour(conn, sejour_id)
    conn.close()
    return res


# ----------------------------------------------------------------------------
# Association activité ↔ créneau
# ----------------------------------------------------------------------------

class SlotIn(BaseModel):
    jour: str
    moment: str
    activite_id: int
    attendu_ids: List[int] = []


def _check_slot(jour: str, moment: str):
    if jour not in JOURS or moment not in MOMENTS:
        raise HTTPException(400, "Jour ou créneau invalide")


@router.put("/sejours/{sejour_id}/slot")
def ajouter_ou_maj_slot(sejour_id: int, slot: SlotIn):
    """Ajoute l'activité au créneau (ou met à jour ses attendus si déjà présente)."""
    _check_slot(slot.jour, slot.moment)
    ensure_schema()
    conn = get_db()
    if not conn.execute("SELECT 1 FROM sejour WHERE id=?", (sejour_id,)).fetchone():
        raise HTTPException(404, "Séjour introuvable")
    if not conn.execute("SELECT 1 FROM activite WHERE id=?", (slot.activite_id,)).fetchone():
        raise HTTPException(404, "Activité introuvable")

    seq = conn.execute("SELECT id FROM sequence WHERE sejour_id=? AND date=? AND moment=?",
                       (sejour_id, slot.jour, slot.moment)).fetchone()
    if seq:
        seq_id = seq["id"]
    else:
        ordre = JOURS.index(slot.jour) * 2 + MOMENTS.index(slot.moment)
        seq_id = conn.execute(
            "INSERT INTO sequence (sejour_id, nom, date, moment, ordre) VALUES (?,?,?,?,?)",
            (sejour_id, f"{slot.jour.capitalize()} {slot.moment}", slot.jour, slot.moment, ordre)
        ).lastrowid

    if not conn.execute("SELECT 1 FROM sequence_activite WHERE sequence_id=? AND activite_id=?",
                        (seq_id, slot.activite_id)).fetchone():
        nxt = conn.execute("SELECT COALESCE(MAX(ordre),-1)+1 FROM sequence_activite WHERE sequence_id=?",
                           (seq_id,)).fetchone()[0]
        conn.execute("INSERT INTO sequence_activite (sequence_id, activite_id, ordre) VALUES (?,?,?)",
                     (seq_id, slot.activite_id, nxt))

    conn.execute("DELETE FROM sequence_activite_attendu WHERE sequence_id=? AND activite_id=?",
                 (seq_id, slot.activite_id))
    # Seuls les attendus réellement rattachés à l'activité sont retenus
    valides = {r[0] for r in conn.execute(
        "SELECT attendu_id FROM activite_attendu WHERE activite_id=?", (slot.activite_id,))}
    for att in slot.attendu_ids:
        if att in valides:
            conn.execute("INSERT INTO sequence_activite_attendu VALUES (?,?,?)",
                         (seq_id, slot.activite_id, att))
    conn.commit()
    res = _charger_sejour(conn, sejour_id)
    res["doc2"] = _maj_doc2(conn, sejour_id)
    conn.close()
    return res


@router.delete("/sejours/{sejour_id}/slot")
def retirer_slot(sejour_id: int, jour: str, moment: str, activite_id: int):
    _check_slot(jour, moment)
    ensure_schema()
    conn = get_db()
    seq = conn.execute("SELECT id FROM sequence WHERE sejour_id=? AND date=? AND moment=?",
                       (sejour_id, jour, moment)).fetchone()
    if seq:
        conn.execute("DELETE FROM sequence_activite WHERE sequence_id=? AND activite_id=?",
                     (seq["id"], activite_id))
        conn.commit()
    res = _charger_sejour(conn, sejour_id)
    res["doc2"] = _maj_doc2(conn, sejour_id)
    conn.close()
    return res


# ----------------------------------------------------------------------------
# Ressources en autonomie (jeux, supports… tirés de la base)
# ----------------------------------------------------------------------------

class RessourceIn(BaseModel):
    activite_id: int


@router.put("/sejours/{sejour_id}/ressource")
def ajouter_ressource(sejour_id: int, r: RessourceIn):
    ensure_schema()
    conn = get_db()
    if not conn.execute("SELECT 1 FROM sejour WHERE id=?", (sejour_id,)).fetchone():
        raise HTTPException(404, "Séjour introuvable")
    if not conn.execute("SELECT 1 FROM activite WHERE id=?", (r.activite_id,)).fetchone():
        raise HTTPException(404, "Activité introuvable")
    nxt = conn.execute("SELECT COALESCE(MAX(ordre),-1)+1 FROM sejour_ressource WHERE sejour_id=?",
                       (sejour_id,)).fetchone()[0]
    conn.execute("INSERT OR IGNORE INTO sejour_ressource (sejour_id, activite_id, ordre) VALUES (?,?,?)",
                 (sejour_id, r.activite_id, nxt))
    conn.commit()
    res = _charger_sejour(conn, sejour_id)
    conn.close()
    return res


@router.delete("/sejours/{sejour_id}/ressource/{activite_id}")
def retirer_ressource(sejour_id: int, activite_id: int):
    ensure_schema()
    conn = get_db()
    conn.execute("DELETE FROM sejour_ressource WHERE sejour_id=? AND activite_id=?",
                 (sejour_id, activite_id))
    conn.commit()
    res = _charger_sejour(conn, sejour_id)
    conn.close()
    return res


# ----------------------------------------------------------------------------
# Document Word
# ----------------------------------------------------------------------------

def _dossier_sortie(s: dict) -> Path:
    """Dossier du séjour dans Dropbox s'il existe, sinon sorties_sejours/."""
    if s.get("dossier_path") and Path(s["dossier_path"]).is_dir():
        return Path(s["dossier_path"])
    d = D.trouver_dossier(s)
    if d:
        return d
    SORTIES_DIR.mkdir(exist_ok=True)
    return SORTIES_DIR


def _nom_fichier(s: dict) -> str:
    return f"Planning détaillé - {s['date_debut'].replace('-', '')} - {D.nom_court(s.get('enseignant'))}.docx"


VERT_FONCE = "2E5238"
VERT = "4A7C59"
VERT_CLAIR = "E6F0E9"
CREME = "FBF7E8"
GRIS = "6B6F6A"
GRIS_CLAIR = "F4F5F2"
COULEURS_UNITE = {"orange": "F28C28", "rose": "D6336C", "bleue": "2F8FD0", "verte": "2F8F46"}

JOURS_FR = ["lundi", "mardi", "mercredi", "jeudi", "vendredi", "samedi", "dimanche"]
MOIS_FR = ["janvier", "février", "mars", "avril", "mai", "juin", "juillet",
           "août", "septembre", "octobre", "novembre", "décembre"]

# Présentation brève de chaque pôle (d'après « Présentation des pôles de l'éco-centre de Chamagnieu »)
POLES = {
    "potager": {
        "titre": "Le potager",
        "texte": ("La matinée débute par une courte partie théorique adaptée au niveau des élèves "
                  "(fonctionnement des plantes, saisonnalité, techniques de culture, alimentation…), "
                  "suivie d'une partie pratique (compost, semis, désherbage, arrosage, paillage, plantation…) "
                  "et d'un bilan. La classe est divisée en trois groupes, chacun accompagné d'un adulte, "
                  "qui tournent sur trois ateliers. L'objectif : participer à l'épanouissement du potager "
                  "au fil des saisons et sensibiliser les enfants à l'impact de notre consommation alimentaire."),
        "regles": "Ne pas arracher de plantes sans l'accord d'un adulte, ne pas marcher sur les barrières en bois, "
                  "rester sur les chemins en copeaux.",
    },
    "miniferme": {
        "titre": "La mini-ferme",
        "texte": ("Les classes sont attendues à 10 h devant le portail principal de la mini-ferme. Après un rappel "
                  "des règles de sécurité, la classe est divisée en trois groupes, chacun avec un adulte référent, "
                  "et les élèves participent à tour de rôle aux soins des animaux (paillage, pansage, nourrissage…). "
                  "Ils découvrent les espèces, leurs besoins et le vocabulaire qui s'y rattache, "
                  "et développent empathie, responsabilité et respect du vivant."),
        "regles": "Marcher dans les enclos, parler à voix basse, laisser les animaux au sol, "
                  "être calme et doux avec tous les animaux.",
    },
    "biodiversite": {
        "titre": "Le pôle biodiversité",
        "texte": ("La matinée alterne une partie théorique (qu'est-ce que la biodiversité ?), une partie pratique "
                  "et un bilan. Les élèves agissent en faveur de la biodiversité (chantiers écologiques, nichoirs, "
                  "mangeoires, gîtes…), apprennent à reconnaître la faune et la flore du parc (inventaires, pièges "
                  "photographiques) et pratiquent des sciences participatives (Vigie-Nature École, Birdlab, Plantnet…)."),
        "regles": "Activités manuelles en extérieur : prévoir des vêtements et chaussures adaptés "
                  "(on peut se salir) et utiliser le matériel de façon responsable.",
    },
    "libre": {
        "titre": "Activité libre",
        "texte": "Matinée laissée à l'initiative de l'enseignant.",
        "regles": None,
    },
}


def _cle_pole(label: str) -> Optional[str]:
    t = label.lower()
    if "potager" in t:
        return "potager"
    if "ferme" in t:
        return "miniferme"
    if "biodiv" in t:
        return "biodiversite"
    if "libre" in t:
        return "libre"
    return None


def _libelle_matin(cle: Optional[str], label: str) -> str:
    """Nom du pôle ; si tu as ajouté des précisions dans le document, elles sont conservées."""
    if not cle:
        return label[:1].upper() + label[1:]
    titre = POLES[cle]["titre"]
    if D.norm(label) in {"polebiodiversite", "miniferme", "potager", "activitelibre"}:
        return titre
    return f"{titre} · {label}"


def _parse_matin(texte: Optional[str]):
    """« Mardi : biodiversité\\nMercredi : mini ferme… » → [(jour, label, clé_pôle)]."""
    out = []
    for ligne in (texte or "").splitlines():
        m = re.match(r"\s*(lundi|mardi|mercredi|jeudi|vendredi)\s*:\s*(.+?)\s*$", ligne, re.I)
        if m:
            out.append((m.group(1).capitalize(), m.group(2), _cle_pole(m.group(2))))
    return out


def _date_longue(iso: str) -> str:
    y, m, d = (int(x) for x in iso.split("-"))
    return f"{JOURS_FR[date(y, m, d).weekday()]} {d}{'er' if d == 1 else ''} {MOIS_FR[m - 1]}"


# --- Helpers python-docx bas niveau ---

def _shade(cell, hex_color):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    tcPr = cell._tc.get_or_add_tcPr()
    for old in tcPr.findall(qn("w:shd")):
        tcPr.remove(old)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def _borders(cell, **kw):
    """kw : top/left/bottom/right = (taille_1/8pt, couleur) ou None pour aucune bordure."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    tcPr = cell._tc.get_or_add_tcPr()
    b = tcPr.find(qn("w:tcBorders"))
    if b is None:
        b = OxmlElement("w:tcBorders")
        tcPr.append(b)
    for side in ("top", "left", "bottom", "right"):
        spec = kw.get(side)
        el = OxmlElement(f"w:{side}")
        if spec:
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), str(spec[0]))
            el.set(qn("w:color"), spec[1])
        else:
            el.set(qn("w:val"), "nil")
        el.set(qn("w:space"), "0")
        b.append(el)


def _marges(cell, haut=80, bas=80, gauche=140, droite=140):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    tcPr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for side, v in (("top", haut), ("left", gauche), ("bottom", bas), ("right", droite)):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(v))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tcPr.append(mar)


def _no_split(row):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    trPr = row._tr.get_or_add_trPr()
    trPr.append(OxmlElement("w:cantSplit"))


def _largeur(table, cm_list):
    from docx.shared import Cm
    table.autofit = False
    for row in table.rows:
        for i, w in enumerate(cm_list):
            row.cells[i].width = Cm(w)


def _run(p, texte, taille=None, gras=False, italique=False, couleur=None, maj=False):
    from docx.shared import Pt, RGBColor
    r = p.add_run(texte)
    r.bold, r.italic = gras, italique
    if taille:
        r.font.size = Pt(taille)
    if couleur:
        r.font.color.rgb = RGBColor.from_string(couleur)
    if maj:
        r.font.all_caps = True
    return r


def _para(container, espace_avant=0, espace_apres=4, keep_next=False, align=None):
    from docx.shared import Pt
    p = container.add_paragraph()
    pf = p.paragraph_format
    pf.space_before, pf.space_after = Pt(espace_avant), Pt(espace_apres)
    pf.keep_with_next = keep_next
    if align is not None:
        p.alignment = align
    return p


def _premier_para(cell):
    """Réutilise le paragraphe vide initial d'une cellule."""
    from docx.shared import Pt
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    return p


def _champ_page(p, taille=8.5, couleur=GRIS):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    r = _run(p, "", taille=taille, couleur=couleur)
    for typ, txt in (("begin", None), (None, "PAGE"), ("end", None)):
        if typ:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), typ)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = txt
        r._r.append(el)


def _espace(doc, pts=6):
    from docx.shared import Pt
    p = doc.add_paragraph()
    p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = Pt(pts)
    return p


def _titre_section(doc, texte):
    """Titre de section avec filet vert dessous."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    p = _para(doc, espace_avant=14, espace_apres=6, keep_next=True)
    _run(p, texte, taille=15, gras=True, couleur=VERT_FONCE)
    pPr = p._p.get_or_add_pPr()
    bd = OxmlElement("w:pBdr")
    bot = OxmlElement("w:bottom")
    bot.set(qn("w:val"), "single")
    bot.set(qn("w:sz"), "8")
    bot.set(qn("w:space"), "2")
    bot.set(qn("w:color"), VERT)
    bd.append(bot)
    pPr.append(bd)
    return p


def _details_ressource(a: dict) -> str:
    bouts = []
    if a.get("effectif_min") and a.get("effectif_max"):
        bouts.append(f"{a['effectif_min']} à {a['effectif_max']} joueurs" if a["effectif_min"] != a["effectif_max"]
                     else f"{a['effectif_min']} joueurs")
    elif a.get("effectif_min"):
        bouts.append(f"à partir de {a['effectif_min']} joueurs")
    if a.get("duree_min"):
        bouts.append(f"environ {a['duree_min']} min")
    if a.get("format_groupe"):
        bouts.append(a["format_groupe"])
    return "  ·  ".join(bouts)


def _carte_activite(doc, a, ressource=False):
    from docx.shared import Pt
    t = doc.add_table(rows=1, cols=1)
    t.autofit = False
    c = t.rows[0].cells[0]
    _no_split(t.rows[0])
    _largeur(t, [17.0])
    _shade(c, GRIS_CLAIR)
    _borders(c, left=(36, VERT))
    _marges(c, 100, 100, 200, 160)

    p = _premier_para(c)
    p.paragraph_format.keep_with_next = True
    _run(p, a["nom"], taille=12, gras=True, couleur=VERT_FONCE)

    if a["description"]:
        p = _para(c, espace_avant=3, espace_apres=3)
        _run(p, a["description"], taille=10)
    if ressource:
        det = _details_ressource(a)
        if det:
            p = _para(c, espace_avant=2, espace_apres=0)
            _run(p, det, taille=9, gras=True, couleur=GRIS)
    if not ressource and a["objectif_texte"]:
        p = _para(c, espace_avant=3, espace_apres=3)
        _run(p, "Visée pédagogique  ", taille=8.5, gras=True, couleur=VERT, maj=True)
        _run(p, a["objectif_texte"], taille=10, italique=True)
    if not ressource and a["attendus"]:
        p = _para(c, espace_avant=4, espace_apres=1, keep_next=True)
        _run(p, "Attendus scolaires", taille=8.5, gras=True, couleur=VERT, maj=True)
        for at in a["attendus"]:
            dom = " – ".join(x for x in (at["domaine"], at["sous_domaine"]) if x)
            p = _para(c, espace_apres=1)
            p.paragraph_format.left_indent = Pt(10)
            p.paragraph_format.first_line_indent = Pt(-10)
            _run(p, "• ", taille=9, couleur=VERT)
            _run(p, f"[{at['cycle_code']}] ", taille=9, gras=True, couleur=GRIS)
            if dom:
                _run(p, dom + " : ", taille=9, gras=True)
            _run(p, at["libelle"], taille=9)
    _espace(doc, 8)


MARQUEUR_RESSOURCES = "Ressources complémentaires (à compléter)"


def _texte_el(e) -> str:
    return "".join(e.xpath(".//w:t/text()"))


def _lire_zone_manuscrite(chemin: Path):
    """Récupère ce que tu as écrit à la main sous le marqueur, dans la version déjà générée."""
    if not chemin.exists():
        return [], False
    try:
        from docx import Document
        corps = Document(chemin).element.body
    except Exception:
        return [], False
    elems = list(corps)
    idx = None
    for i, e in enumerate(elems):
        if e.tag.endswith("}p") and _texte_el(e).strip().lower() == MARQUEUR_RESSOURCES.lower():
            idx = i
    if idx is None:
        return [], False
    queue = [e for e in elems[idx + 1:] if not e.tag.endswith("}sectPr")]
    # on ignore la phrase d'invite si elle n'a pas été modifiée
    queue = [e for e in queue
             if not _texte_el(e).strip().startswith("(Écris ici les ressources complémentaires")]
    if not any(_texte_el(e).strip() or e.xpath(".//w:drawing|.//w:pict") for e in queue):
        return [], False
    avec_images = any(e.xpath(".//w:drawing|.//w:pict") for e in queue)
    # un élément contenant une image pointe vers des relations de l'ancien fichier : on ne le reporte pas
    queue = [e for e in queue if not e.xpath(".//w:drawing|.//w:pict")]
    return queue, avec_images


def _greffer_zone_manuscrite(doc, elements):
    import copy
    corps = doc.element.body
    sect = corps.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}sectPr")
    for e in elements:
        sect.addprevious(copy.deepcopy(e))


def _sauvegarder_ancienne_version(chemin: Path):
    """Les images ne peuvent pas être reportées : on garde l'ancienne version à côté."""
    import shutil
    shutil.copy2(chemin, chemin.with_name(chemin.stem + " (sauvegarde avant régénération).docx"))


def generer_docx(s: dict) -> Path:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt

    # Le nouveau document se base sur « déroulé et objectifs » (retouches comprises)
    d2 = {}
    f2 = D2.trouver_doc2(_dossier_sortie(s))
    if f2:
        try:
            d2 = D2.lire_doc2(f2)
        except Exception:
            d2 = {}

    d = Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin, sec.bottom_margin = Cm(1.6), Cm(1.8)
    normal = d.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(4)

    # Pied de page
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run(fp, "Éco-centre de Chamagnieu  ·  Récapitulatif de séjour  ·  page ", taille=8.5, couleur=GRIS)
    _champ_page(fp)

    # --- Bandeau de titre ---
    t = d.add_table(rows=1, cols=1)
    t.autofit = False
    c = t.rows[0].cells[0]
    _largeur(t, [17.0])
    _shade(c, VERT_FONCE)
    _borders(c)
    _marges(c, 280, 280, 320, 320)
    p = _premier_para(c)
    _run(p, "VOTRE SÉJOUR À L'ÉCO-CENTRE DE CHAMAGNIEU", taille=9.5, gras=True, couleur="BFE0C8")
    p = _para(c, espace_avant=4, espace_apres=2)
    _run(p, "Récapitulatif du projet", taille=24, gras=True, couleur="FFFFFF")
    p = _para(c, espace_avant=2, espace_apres=0)
    _run(p, f"Du {_date_longue(s['date_debut'])} au {_date_longue(s['date_fin'])} {s['date_fin'][:4]}",
         taille=12, couleur="FFFFFF")
    _espace(d, 10)

    # --- Fiche d'identité du séjour (colonnes A, B, D, E, F, H, I, J — sans la catégorie) ---
    niveau = " / ".join(x for x in (s["niveau1"], s["niveau2"]) if x)
    infos = [
        ("Enseignant(e)", d2.get("enseignant") or s["enseignant"]),
        ("École", s["ecole"]),
        ("Niveau", niveau),
        ("Classe AR", s["classe_ar"]),
        ("Unité", (d2.get("unite") or s["unite"] or "").capitalize() or None),
        ("Thématique", d2.get("thematique") or s["thematique_excel"]),
        ("Objectifs", d2.get("objectifs") or s.get("objectifs_texte")),
    ]
    infos = [(k, v) for k, v in infos if v]
    t = d.add_table(rows=len(infos), cols=2)
    t.autofit = False
    for i, (k, v) in enumerate(infos):
        kc, vc = t.rows[i].cells
        _shade(kc, VERT_CLAIR)
        for cell in (kc, vc):
            _borders(cell, bottom=(4, "D5DDD7"))
            _marges(cell, 70, 70, 140, 140)
        _run(_premier_para(kc), k, taille=10, gras=True, couleur=VERT_FONCE)
        pv = _premier_para(vc)
        if k == "Unité":
            col = COULEURS_UNITE.get(v.lower())
            if col:
                _run(pv, "■ ", taille=11, couleur=col)
        _run(pv, v, taille=10.5)
    _largeur(t, [4.0, 13.0])

    # --- Matinées ---
    texte_matin = s["matin_texte"]
    if d2.get("matins"):
        texte_matin = "\n".join(f"{j} : {t}" for j, t in d2["matins"].items())
    matin = _parse_matin(texte_matin)
    _titre_section(d, "Les matinées : les pôles de l'éco-centre")
    if matin:
        p = _para(d, espace_apres=6)
        _run(p, "Chaque matin, la classe est accueillie sur l'un des pôles de l'éco-centre, "
                "selon le roulement de son unité.", taille=10, italique=True, couleur=GRIS)
        t = d.add_table(rows=len(matin), cols=2)
        t.autofit = False
        for i, (jour, label, cle) in enumerate(matin):
            jc, pc = t.rows[i].cells
            _no_split(t.rows[i])
            _shade(jc, VERT)
            _borders(jc, bottom=(8, "FFFFFF"))
            _borders(pc, bottom=(4, "D5DDD7"))
            _marges(jc, 70, 70, 140, 100)
            _marges(pc, 70, 70, 160, 100)
            _run(_premier_para(jc), jour, taille=10.5, gras=True, couleur="FFFFFF")
            _run(_premier_para(pc), _libelle_matin(cle, label), taille=10.5, gras=True)
        _largeur(t, [3.5, 13.5])

        vus = []
        for _, _, cle in matin:
            if cle and cle != "libre" and cle not in vus:
                vus.append(cle)
        for cle in vus:
            pole = POLES[cle]
            _espace(d, 8)
            tt = d.add_table(rows=1, cols=1)
            tt.autofit = False
            cc = tt.rows[0].cells[0]
            _no_split(tt.rows[0])
            _largeur(tt, [17.0])
            _shade(cc, CREME)
            _borders(cc, left=(36, "E0B94A"))
            _marges(cc, 100, 100, 200, 160)
            p = _premier_para(cc)
            p.paragraph_format.keep_with_next = True
            _run(p, pole["titre"], taille=12, gras=True, couleur=VERT_FONCE)
            p = _para(cc, espace_avant=2, espace_apres=3)
            _run(p, pole["texte"], taille=10)
            if pole["regles"]:
                p = _para(cc, espace_apres=0)
                _run(p, "À savoir  ", taille=8.5, gras=True, couleur="9A7A1E", maj=True)
                _run(p, pole["regles"], taille=10, italique=True)
    elif texte_matin:
        _para(d).add_run(texte_matin)
    else:
        p = _para(d)
        _run(p, "Programme des matinées à venir.", italique=True, couleur=GRIS)

    # --- Après-midis et veillées ---
    _titre_section(d, "Les après-midis et les veillées")
    for jour in JOURS:
        # bandeau du jour
        t = d.add_table(rows=1, cols=1)
        t.autofit = False
        c = t.rows[0].cells[0]
        _largeur(t, [17.0])
        _shade(c, VERT_CLAIR)
        _borders(c, bottom=(8, VERT))
        _marges(c, 60, 60, 160, 160)
        p = _premier_para(c)
        p.paragraph_format.keep_with_next = True
        _run(p, jour.upper(), taille=12, gras=True, couleur=VERT_FONCE)
        # garder le bandeau avec ce qui suit
        _espace(d, 4).paragraph_format.keep_with_next = True

        for moment in MOMENTS:
            p = _para(d, espace_avant=2, espace_apres=4, keep_next=True)
            _run(p, "Après-midi" if moment == "après-midi" else "Veillée", taille=10.5, gras=True,
                 couleur=VERT, maj=True)
            acts = s["planning"][jour][moment]
            if not acts:
                p = _para(d, espace_apres=6)
                _run(p, "À définir", italique=True, couleur=GRIS)
                continue
            for a in acts:
                _carte_activite(d, a)
        _espace(d, 6)

    # --- Ressources en autonomie ---
    dossier = _dossier_sortie(s)
    chemin = dossier / _nom_fichier(s)
    _titre_section(d, "Ressources en autonomie pour la classe")
    p = _para(d, espace_apres=6)
    _run(p, "Jeux et supports que les élèves peuvent utiliser seuls dans la salle de classe.",
         taille=10, italique=True, couleur=GRIS)
    for r in s.get("ressources", []):
        _carte_activite(d, r, ressource=True)
    p = _para(d, espace_avant=6, espace_apres=4, keep_next=True)
    _run(p, MARQUEUR_RESSOURCES, taille=10.5, gras=True, couleur=VERT, maj=True)
    reprise, avec_images = _lire_zone_manuscrite(chemin)
    if avec_images:
        _sauvegarder_ancienne_version(chemin)
    if reprise:
        _greffer_zone_manuscrite(d, reprise)
    else:
        p = _para(d, espace_apres=4)
        _run(p, "(Écris ici les ressources complémentaires : elles seront conservées si tu "
                "régénères le document.)", taille=10, italique=True, couleur=GRIS)

    try:
        d.save(chemin)
    except PermissionError:
        raise HTTPException(423, f"Impossible d'écrire {chemin.name} (ouvert dans Word ?)")
    return chemin


@router.post("/sejours/{sejour_id}/document")
def creer_document(sejour_id: int, ouvrir: bool = True):
    ensure_schema()
    conn = get_db()
    s = _charger_sejour(conn, sejour_id)
    conn.close()
    chemin = generer_docx(s)
    if ouvrir and os.name == "nt":
        os.startfile(chemin)
    return {"chemin": str(chemin)}
