# =============================================================================
# typo.py
# =============================================================================
# Applique à un document « Déroulé et objectifs » la typographie d'un document de référence
# (celui que tu as mis en forme dans Word) : police et interligne par défaut, tailles de texte,
# largeur des colonnes du tableau de détail, pagination.
#
# La transformation est absolue (tailles fixées par rôle, pas décalées) : on peut donc
# la rejouer sur un document déjà au bon format sans rien changer. Le contenu n'est jamais touché.
# =============================================================================

import copy
import shutil
from pathlib import Path
from typing import List, Optional

from docx import Document
from docx.oxml.ns import qn

# Tailles en demi-points (22 = 11 pt), relevées sur le document de référence
TAILLE_TITRE_PETIT = 20      # « DÉROULÉ ET OBJECTIFS PÉDAGOGIQUES »
TAILLE_TITRE_DATES = 44      # « Séjour du … au … »
TAILLE_TITRES = 28           # « Emploi du temps de la semaine », « Détail des activités… »
TAILLE_EDT_ENTETE = 22       # Lundi … Vendredi
TAILLE_EDT_LIBELLE = 21      # Matin, Après-midi, Veillée
TAILLE_DETAIL_JOUR = 26      # « Lundi » dans le tableau de détail
TAILLE_DETAIL_SOUS_JOUR = 19  # « Après-midi » sous le jour
TAILLE_DETAIL_LIBELLE = 21   # « Objectifs pédagogiques », « Descriptif de l'activité »
LARGEURS_DETAIL = (1500, 1650, 7350)
INTERLIGNE = "288"


def _taille(run, demi_points: int) -> bool:
    """Fixe la taille d'un morceau de texte qui en a déjà une. Retourne True s'il y a eu un changement."""
    rpr = run._r.rPr
    if rpr is None:
        return False
    change = False
    for balise in ("w:sz", "w:szCs"):
        el = rpr.find(qn(balise))
        if el is not None and el.get(qn("w:val")) != str(demi_points):
            el.set(qn("w:val"), str(demi_points))
            change = True
    return change


def _runs(paragraphe):
    return paragraphe.runs


def _cellules_distinctes(row) -> list:
    vues, out = [], []
    for c in row.cells:
        if c._tc not in vues:
            vues.append(c._tc)
            out.append(c)
    return out


def _interligne(paragraphe) -> bool:
    sp = paragraphe._p.pPr.find(qn("w:spacing")) if paragraphe._p.pPr is not None else None
    if sp is not None and sp.get(qn("w:line")) == "276":
        sp.set(qn("w:line"), INTERLIGNE)
        return True
    return False


def appliquer_typographie(chemin: Path, reference: Path) -> List[str]:
    """Met le document au format de la référence. Retourne la liste des changements faits ([] = déjà conforme)."""
    doc, ref = Document(chemin), Document(reference)
    modifs: List[str] = []

    # 1. police et interligne par défaut
    dd_ref = ref.styles.element.find(qn("w:docDefaults"))
    dd = doc.styles.element.find(qn("w:docDefaults"))
    if dd_ref is not None and dd is not None:
        from lxml import etree
        if etree.tostring(dd) != etree.tostring(dd_ref):
            doc.styles.element.replace(dd, copy.deepcopy(dd_ref))
            modifs.append("police et interligne par défaut (Calibri 11, interligne 1,2)")

    # 2. titres du haut de page et titres de section
    for p in doc.paragraphs:
        t = p.text.strip()
        cible = None
        if t.startswith("DÉROULÉ ET OBJECTIFS"):
            cible = TAILLE_TITRE_PETIT
        elif t.startswith("Séjour du"):
            cible = TAILLE_TITRE_DATES
        elif t.startswith(("Emploi du temps de la semaine", "Détail des activités", "Ressources en autonomie")):
            cible = TAILLE_TITRES
        if cible and any([_taille(r, cible) for r in _runs(p)]):
            modifs.append(f"taille du titre « {t[:30]} »")
        if t.startswith("Séjour du"):  # les dates ne sont plus en italique
            retire = False
            for r in p.runs:
                i = r._r.rPr.find(qn("w:i")) if r._r.rPr is not None else None
                if i is not None:
                    r._r.rPr.remove(i)
                    retire = True
            if retire:
                modifs.append("dates du titre en caractères droits")
        if _interligne(p):
            modifs.append("interligne d'un paragraphe")

    tables = doc.tables
    # 3. emploi du temps de la semaine
    if len(tables) > 1 and tables[1].rows[1].cells[0].text.strip().lower() == "matin":
        edt = tables[1]
        n = 0
        for c in _cellules_distinctes(edt.rows[0]):
            for p in c.paragraphs:
                n += any([_taille(r, TAILLE_EDT_ENTETE) for r in p.runs])
        for row in edt.rows[1:]:
            for p in _cellules_distinctes(row)[0].paragraphs:
                n += any([_taille(r, TAILLE_EDT_LIBELLE) for r in p.runs])
        if n:
            modifs.append("tailles du tableau « emploi du temps »")

    # 4. détail des activités
    if len(tables) > 2 and _cellules_distinctes(tables[2].rows[0])[1].text.strip().startswith("Objectifs"):
        from docx.text.paragraph import Paragraph
        detail = tables[2]
        n = 0
        for i, tr in enumerate(detail._tbl.findall(qn("w:tr"))):
            tcs = tr.findall(qn("w:tc"))
            paras = [[Paragraph(p, detail) for p in tc.findall(qn("w:p"))] for tc in tcs]
            for liste in paras:
                for p in liste:
                    n += _interligne(p)
            if len(tcs) < 3:
                continue
            if i % 2 == 0:  # ligne « Objectifs pédagogiques » : elle reste avec la ligne « Descriptif » qui suit
                if paras[0]:
                    n += any([_taille(r, TAILLE_DETAIL_JOUR) for r in paras[0][0].runs])
                if len(paras[0]) > 1:
                    n += any([_taille(r, TAILLE_DETAIL_SOUS_JOUR) for r in paras[0][1].runs])
                for liste in paras:
                    for p in liste:
                        if not p.paragraph_format.keep_with_next:
                            p.paragraph_format.keep_with_next = True
                            n += 1
            for p in paras[1]:
                n += any([_taille(r, TAILLE_DETAIL_LIBELLE) for r in p.runs])
        # largeur des colonnes
        grille = detail._tbl.find(qn("w:tblGrid"))
        for col, w in zip(grille.findall(qn("w:gridCol")), LARGEURS_DETAIL):
            if col.get(qn("w:w")) != str(w):
                col.set(qn("w:w"), str(w))
                n += 1
        for tr in detail._tbl.findall(qn("w:tr")):
            for tc, w in zip(tr.findall(qn("w:tc")), LARGEURS_DETAIL):
                tcw = tc.find(qn("w:tcPr")).find(qn("w:tcW"))
                if tcw is not None and tcw.get(qn("w:w")) != str(w):
                    tcw.set(qn("w:w"), str(w))
                    n += 1
        if n:
            modifs.append("tableau « détail des activités » (tailles, colonnes, pagination)")

    if modifs:
        doc.save(chemin)
    return modifs


def sauvegarder(chemin: Path, racine: Path, dossier_sauvegarde: Path) -> Path:
    """Copie le fichier (même arborescence) avant modification."""
    try:
        rel = chemin.relative_to(racine)
    except ValueError:
        rel = Path(chemin.name)
    cible = dossier_sauvegarde / rel
    cible.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(chemin, cible)
    return cible
